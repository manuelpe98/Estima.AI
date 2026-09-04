"""Genera la relazione descrittiva del computo in Word, in prosa discorsiva
(per coerenza con la convenzione già seguita per il caso Colombo)."""
from __future__ import annotations
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from ..models import ComputoVoce, ProjectMeta, RoomComparison, ORIGINE_INFO

# Colore di evidenziazione per le voci "da completare a mano" (stesso giallo
# usato nel file Excel), applicato allo sfondo delle celle della riga.
_DA_COMPLETARE_FILL_HEX = "FFF2CC"


def _shade_cell(cell, hex_color: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def build_word(voci: list[ComputoVoce], meta: ProjectMeta, note_metodologiche: list[str],
                validazione_msg: list[str], out_path: str,
                confronto: list[RoomComparison] | None = None) -> str:
    doc = Document()

    title = doc.add_heading("Computo Metrico Estimativo", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    tipo_lbl = "nuova costruzione" if meta.tipo_intervento == "nuova_costruzione" else "ristrutturazione"
    doc.add_paragraph(
        f"Il presente documento riporta il computo metrico estimativo relativo al progetto "
        f"\"{meta.nome_progetto}\" ({tipo_lbl})" + (f", per il committente {meta.committente}" if meta.committente else "") +
        (f", ubicato in {meta.ubicazione}" if meta.ubicazione else "") +
        f". Le quantità sono state rilevate in via automatica dagli elaborati grafici caricati "
        f"(pianta quotata in scala, ed eventuale pianta strutturale) e valorizzate sulla base del "
        f"prezzario \"{meta.prezzario_nome}\". Il documento ha carattere preliminare/estimativo ed è "
        "soggetto a verifica da parte del tecnico incaricato prima di ogni utilizzo con valenza contrattuale."
    )

    doc.add_heading("Metodologia e fonti", level=1)
    doc.add_paragraph(
        "Le superfici dei singoli vani, il sedime dell'edificio e il conteggio di serramenti ed elementi "
        "strutturali puntuali (pilastri, travi) sono stati ricavati per via geometrica dagli elaborati "
        "caricati, verificati come elaborati vettoriali in scala e completi di quote. Le scelte di "
        "materiali, finiture e i parametri dimensionali non deducibili dal disegno (altezze di interpiano, "
        "profondità di scavo, incidenze parametriche) sono stati raccolti tramite questionario strutturato "
        "e confermati dall'utente, e sono riportati nelle note metodologiche e nelle singole voci."
    )
    if note_metodologiche:
        critiche = [n for n in note_metodologiche if n.startswith("ATTENZIONE")]
        normali = [n for n in note_metodologiche if not n.startswith("ATTENZIONE")]
        if critiche:
            doc.add_heading("⚠ Avvisi critici — leggere prima di usare questo computo", level=2)
            for nota in critiche:
                p = doc.add_paragraph()
                run = p.add_run(nota)
                run.bold = True
                run.font.color.rgb = RGBColor(0x9C, 0x00, 0x06)
        for nota in normali:
            doc.add_paragraph(nota, style="Intense Quote")

    if validazione_msg:
        doc.add_heading("Esito della verifica degli elaborati grafici", level=1)
        doc.add_paragraph(" ".join(validazione_msg))

    if confronto:
        doc.add_heading("Confronto stato di fatto / stato di progetto", level=1)
        doc.add_paragraph(
            "Il confronto tra i vani rilevati nello stato di fatto e nello stato di progetto individua i "
            "vani demoliti (presenti solo nello stato di fatto), i vani di nuova formazione (presenti solo "
            "nel progetto) e i vani con superficie variata, riportati di seguito. Solo i vani demoliti "
            "generano automaticamente le relative voci di demolizione nel computo; i vani con superficie "
            "variata sono segnalati per una verifica puntuale, non essendo automaticamente quantificabile "
            "l'entità dell'intervento di adeguamento."
        )
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        for i, h in enumerate(["Vano", "Stato", "Area stato di fatto (m²)", "Area stato di progetto (m²)"]):
            hdr[i].text = h
        for c in confronto:
            row = table.add_row().cells
            row[0].text = c.label
            row[1].text = c.stato.replace("_", " ")
            row[2].text = f"{c.area_sdf_m2:.2f}" if c.area_sdf_m2 is not None else "-"
            row[3].text = f"{c.area_sdp_m2:.2f}" if c.area_sdp_m2 is not None else "-"
        doc.add_paragraph()

    ha_commenti = any(getattr(v, "commento", "") for v in voci)
    n_cols = 9 if ha_commenti else 8
    doc.add_heading("Elenco voci di computo", level=1)
    table = doc.add_table(rows=1, cols=n_cols)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    intestazioni = ["N.", "Codice", "Categoria", "Descrizione", "U.M.", "Q.tà", "Importo (€)", "Affidabilità"]
    if ha_commenti:
        intestazioni.append("Commento")
    for i, h in enumerate(intestazioni):
        hdr[i].text = h

    totale = 0.0
    righe_da_completare = 0
    for v in voci:
        da_completare = getattr(v, "da_completare", False)
        row = table.add_row().cells
        row[0].text = str(v.numero)
        row[1].text = v.codice
        row[2].text = v.categoria
        descrizione = v.descrizione
        if da_completare:
            righe_da_completare += 1
            descrizione = "⚠ DA COMPLETARE A MANO — " + descrizione
        row[3].text = descrizione
        row[4].text = v.unita_misura
        row[5].text = f"{v.quantita:,.2f}"
        row[6].text = f"{v.importo:,.2f}"
        origine_label, origine_emoji = ORIGINE_INFO.get(getattr(v, "origine", "assunta"), ("", ""))
        row[7].text = f"{origine_emoji} {origine_label}".strip()
        if ha_commenti:
            row[8].text = getattr(v, "commento", "") or ""
        if da_completare:
            for cell in row:
                _shade_cell(cell, _DA_COMPLETARE_FILL_HEX)
        totale += v.importo

    if righe_da_completare:
        legend_p = doc.add_paragraph()
        legend_run = legend_p.add_run(
            f"⚠ {righe_da_completare} voce/i evidenziata/e in giallo: da completare a mano "
            "(quantità/prezzo non calcolabili dagli elaborati caricati)."
        )
        legend_run.bold = True

    legenda_p = doc.add_paragraph()
    legenda_run = legenda_p.add_run(
        "Legenda colonna Affidabilità (provenienza del dato, categorie da non mischiare tra loro): "
        + "; ".join(f"{emoji} {label}" for label, emoji in
                     (ORIGINE_INFO["esplicita"], ORIGINE_INFO["derivata"], ORIGINE_INFO["inferita"],
                      ORIGINE_INFO["assunta"], ORIGINE_INFO["non_disponibile"])) + "."
    )
    legenda_run.italic = True

    doc.add_paragraph()
    tot_p = doc.add_paragraph()
    tot_run = tot_p.add_run(f"Totale computo (al netto di IVA, oneri di sicurezza e spese tecniche): € {totale:,.2f}")
    tot_run.bold = True
    tot_run.font.size = Pt(12)

    doc.add_heading("Avvertenze", level=1)
    doc.add_paragraph(
        "Il presente computo copre le opere di finitura (pavimenti, pareti interne, serramenti, porte "
        "interne), l'approntamento di cantiere, le fondazioni (magrone, calcestruzzo, casseforme e acciaio "
        "come voci separate) e gli scavi, le strutture in elevazione (calcestruzzo, casseforme e acciaio), "
        "i solai (a pacchetto completo o, se richiesto, scomposti in casseforme/calcestruzzo/acciaio per una "
        "soletta piena), la copertura, il cappotto termico, le impermeabilizzazioni contro terra, l'eventuale "
        "piscina individuata in pianta (scavo, vasca, impermeabilizzazione e bordo) e, dove richiesto dalle "
        "risposte al questionario, il vespaio aerato, le pareti divisorie interne, le contropareti e le "
        "velette, oltre a una stima forfettaria degli impianti elettrico e idrico-sanitario — nei limiti di "
        "ciò che è quantificabile a partire dagli elaborati caricati e dai parametri confermati dall'utente "
        "(spessori, incidenze di armatura, lunghezze parametriche). Restano fuori da questa versione, per "
        "non rischiare di introdurre valori inventati su elementi che richiedono un vero progetto strutturale "
        "esecutivo (diametri e passi reali delle barre, classi di esposizione per singolo elemento, "
        "casserature per fase costruttiva, spinottature di ripresa getto — quest'ultima è comunque elencata "
        "come voce segnaposto da completare a mano): l'impiantistica di dettaglio, il computo strutturale "
        "esecutivo delle opere in cemento armato e, per le ristrutturazioni, le demolizioni parziali dei "
        "vani con superficie variata (segnalate ma non quantificate automaticamente). Alcune voci — opere "
        "individuate come plausibilmente presenti (es. scala esterna) o tipicamente necessarie ma non "
        "rappresentate nella pianta caricata (es. recinzione, smaltimento acque, sistemazioni esterne) — "
        "sono comunque incluse nell'elenco come segnaposto, evidenziate in giallo con quantità e prezzo a "
        "zero: vanno completate a mano dal tecnico dopo verifica sugli elaborati generali/di dettaglio. Il "
        "prezzario utilizzato, se non caricato esplicitamente dall'utente, è un prezzario di esempio a "
        "scopo dimostrativo e non ha valore ufficiale."
    )

    doc.save(out_path)
    return out_path
