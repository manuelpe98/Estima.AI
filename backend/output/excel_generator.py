"""Genera il computo metrico estimativo in Excel, in stile PriMus."""
from __future__ import annotations
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from ..models import ComputoVoce, ProjectMeta, ORIGINE_INFO

HEADER_FILL = PatternFill(start_color="1F4E5F", end_color="1F4E5F", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
# Evidenzia le voci che il sistema non può quantificare/prezzare da solo
# (ComputoVoce.da_completare): quantità e prezzo a 0, da completare a mano
# prima di considerare il computo definitivo — vedi anche la nota di riga.
DA_COMPLETARE_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
DA_COMPLETARE_FONT = Font(bold=True, color="7F6000")
# Avviso critico (es. "nessun vano riconosciuto"): sfondo rosso tenue, testo in
# grassetto — deve essere impossibile non notarlo aprendo il file, a differenza
# delle altre note metodologiche (vedi ATTENZIONE_PREFIX più sotto). Prima di
# questa versione le note metodologiche non comparivano affatto nel file Excel
# (solo nel Word e nell'interfaccia): chi apriva solo l'Excel non aveva modo di
# sapere PERCHÉ il computo fosse vuoto o incompleto.
ATTENZIONE_FILL = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
ATTENZIONE_FONT = Font(bold=True, color="9C0006")
ATTENZIONE_PREFIX = "ATTENZIONE"

COLUMNS = ["N.", "Codice", "Categoria", "Descrizione", "U.M.", "Quantità",
           "Prezzo unitario (€)", "Importo (€)", "Affidabilità", "Note", "Commento"]


def build_excel(voci: list[ComputoVoce], meta: ProjectMeta, out_path: str,
                 note_metodologiche: list[str] | None = None) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Computo metrico"

    ws["A1"] = "COMPUTO METRICO ESTIMATIVO"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Progetto: {meta.nome_progetto}"
    ws["A3"] = f"Committente: {meta.committente or '-'}"
    ws["A4"] = f"Ubicazione: {meta.ubicazione or '-'}"
    ws["A5"] = f"Prezzario di riferimento: {meta.prezzario_nome}"

    note_row = 6
    if note_metodologiche:
        critiche = [n for n in note_metodologiche if n.startswith(ATTENZIONE_PREFIX)]
        normali = [n for n in note_metodologiche if not n.startswith(ATTENZIONE_PREFIX)]
        for nota in critiche:
            cell = ws.cell(row=note_row, column=1, value=nota)
            cell.fill = ATTENZIONE_FILL
            cell.font = ATTENZIONE_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=11)
            ws.row_dimensions[note_row].height = 45
            note_row += 1
        if normali:
            label_cell = ws.cell(row=note_row, column=1, value="Note metodologiche:")
            label_cell.font = Font(bold=True, italic=True, size=9)
            note_row += 1
            for nota in normali:
                cell = ws.cell(row=note_row, column=1, value=f"• {nota}")
                cell.font = Font(italic=True, size=9)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=11)
                ws.row_dimensions[note_row].height = 30
                note_row += 1
        note_row += 1

    header_row = max(note_row, 7)
    for col_idx, title in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

    row = header_row + 1
    totale = 0.0
    righe_da_completare = 0
    for v in voci:
        da_completare = getattr(v, "da_completare", False)
        note_cell = v.note
        if da_completare:
            righe_da_completare += 1
            note_cell = (("⚠ DA COMPLETARE A MANO — " + note_cell) if note_cell
                         else "⚠ DA COMPLETARE A MANO: quantità e prezzo non calcolabili da questa pianta.")
        origine_label, origine_emoji = ORIGINE_INFO.get(getattr(v, "origine", "assunta"), ("", ""))
        affidabilita_cell = f"{origine_emoji} {origine_label}".strip()
        values = [v.numero, v.codice, v.categoria, v.descrizione, v.unita_misura,
                  v.quantita, v.prezzo_unitario, v.importo, affidabilita_cell, note_cell, getattr(v, "commento", "")]
        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col_idx, value=val)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(col_idx in (4, 10, 11)))
            if col_idx in (6, 7, 8):
                cell.number_format = "#,##0.00"
            if da_completare:
                cell.fill = DA_COMPLETARE_FILL
                if col_idx == 10:
                    cell.font = DA_COMPLETARE_FONT
        totale += v.importo
        row += 1

    if righe_da_completare:
        legend_cell = ws.cell(row=row + 1, column=2, value=f"⚠ {righe_da_completare} voce/i evidenziata/e in giallo: "
                               "da completare a mano (quantità/prezzo non calcolabili da questa pianta)")
        legend_cell.font = DA_COMPLETARE_FONT
        row += 1

    ws.cell(row=row + 1, column=7, value="TOTALE COMPUTO").font = Font(bold=True)
    tot_cell = ws.cell(row=row + 1, column=8, value=round(totale, 2))
    tot_cell.font = Font(bold=True)
    tot_cell.number_format = "#,##0.00"
    row += 3

    ws.cell(row=row, column=2, value="Legenda colonna Affidabilità (provenienza del dato):").font = Font(bold=True)
    row += 1
    for chiave in ("esplicita", "derivata", "inferita", "assunta", "non_disponibile"):
        label, emoji = ORIGINE_INFO[chiave]
        ws.cell(row=row, column=2, value=f"{emoji} {label}")
        row += 1

    widths = [5, 12, 18, 48, 8, 11, 16, 14, 30, 40, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = f"A{header_row + 1}"
    wb.save(out_path)
    return out_path


def build_primus_export(voci: list[ComputoVoce], meta: ProjectMeta, out_path: str) -> str:
    """Genera un file pensato per l'importazione manuale sicura in PriMus.

    Perché non un file 'nativo' PriMus: i formati di interscambio di PriMus
    (XPWE, DCF, PWE) sono formati proprietari ACCA non documentati
    pubblicamente, e la funzione di import Excel di PriMus accetta solo file
    prima esportati da PriMus stesso — non un foglio arbitrario. Generare un
    finto file 'PriMus' rischierebbe di produrre un file che l'applicazione
    rifiuta, o peggio importa in modo silenziosamente sbagliato: inaccettabile
    per un documento con valore economico. Questo file segue invece il
    workflow manuale documentato da ACCA: incolla il foglio "Elenco Prezzi"
    nell'editor Elenco Prezzi di PriMus, poi usa il foglio "Misurazioni" come
    riferimento per inserire le quantità nelle righe di misura.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Elenco Prezzi"
    ws["A1"] = "ELENCO PREZZI — da incollare nell'editor Elenco Prezzi di PriMus"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = ("Vedi il foglio 'Istruzioni' per il procedimento di importazione manuale in PriMus "
                "(PriMus non importa automaticamente fogli Excel non generati da PriMus stesso).")
    header_row = 4
    for col_idx, title in enumerate(["Codice", "Categoria", "Descrizione", "U.M.", "Prezzo unitario (€)"], start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
    row = header_row + 1
    for v in voci:
        for col_idx, val in enumerate([v.codice, v.categoria, v.descrizione, v.unita_misura, v.prezzo_unitario], start=1):
            cell = ws.cell(row=row, column=col_idx, value=val)
            cell.border = BORDER
            if col_idx == 5:
                cell.number_format = "#,##0.00"
        row += 1
    for i, w in enumerate([14, 20, 55, 8, 16], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws2 = wb.create_sheet("Misurazioni")
    ws2["A1"] = "MISURAZIONI — quantità da inserire nelle righe di misura di PriMus dopo l'import dell'elenco prezzi"
    ws2["A1"].font = TITLE_FONT
    header_row2 = 3
    for col_idx, title in enumerate(["Codice", "Descrizione", "U.M.", "Quantità", "Note/ipotesi"], start=1):
        cell = ws2.cell(row=header_row2, column=col_idx, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
    row = header_row2 + 1
    for v in voci:
        for col_idx, val in enumerate([v.codice, v.descrizione, v.unita_misura, v.quantita, v.note], start=1):
            cell = ws2.cell(row=row, column=col_idx, value=val)
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=(col_idx in (2, 5)), vertical="top")
        row += 1
    for i, w in enumerate([14, 45, 8, 12, 45], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    ws3 = wb.create_sheet("Istruzioni")
    istruzioni = [
        "Come importare questo elenco in PriMus (procedimento manuale, per evitare errori di importazione):",
        "1. Apri PriMus e vai nell'editor 'Elenco Prezzi'.",
        "2. Copia le colonne Codice, Descrizione, U.M., Prezzo unitario dal foglio 'Elenco Prezzi' di questo file "
        "e incollale nell'editor secondo la procedura di PriMus (Tutorial PriMus - Importazioni, ACCA software).",
        "3. Crea il computo e trascina le voci importate nelle righe di misura, come di consueto in PriMus.",
        "4. Usa il foglio 'Misurazioni' di questo file come riferimento per inserire la quantità di ciascuna voce "
        "(colonna Quantità) e per leggere le ipotesi/note che hanno originato quel valore.",
        f"Prezzario di riferimento usato per generare questo elenco: {meta.prezzario_nome}.",
        f"Progetto: {meta.nome_progetto}",
    ]
    for i, line in enumerate(istruzioni, start=1):
        ws3.cell(row=i, column=1, value=line)
    ws3.column_dimensions["A"].width = 110

    wb.save(out_path)
    return out_path
